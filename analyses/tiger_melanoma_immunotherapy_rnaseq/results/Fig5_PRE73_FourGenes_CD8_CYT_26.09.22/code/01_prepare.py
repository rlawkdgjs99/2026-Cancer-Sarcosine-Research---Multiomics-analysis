from pathlib import Path
import csv,hashlib,json,math,shutil
import numpy as np
from openpyxl import load_workbook
out=Path(__file__).resolve().parents[1];base=out.parent.parent
src=base/'results/Fig5_PRE73_Ratio_Coordinates_CD8_CYT_26.09.22/tables/patient_data_PRE73.csv'
with src.open() as h:d=list(csv.DictReader(h))
genes=['SARDH','PIPOX','GNMT','DMGDH'];assert len(d)==len({r['sample_id'] for r in d})==len({r['patient_id'] for r in d})==73
snapshot=out/'inputs/73Pre_snapshot.xlsx';shutil.copyfile(base/'73Pre.xlsx',snapshot)
w=load_workbook(snapshot,data_only=True,read_only=True)['02_All_Patient_Data'];headers=[w.cell(10,c).value for c in range(1,26)]
ex={r[1]:dict(zip(headers,r)) for r in w.iter_rows(min_row=11,max_row=83,max_col=25,values_only=True)}
assert set(ex)=={r['sample_id'] for r in d}
for r in d:
 e=ex[r['sample_id']]
 for k in ['patient_id','response','therapy','sex']:assert r[k]==e[k]
 for k in genes+['age']:assert math.isclose(float(r[k]),e[k],rel_tol=1e-11,abs_tol=1e-12)
thresholds=[]
for g in genes:
 a=np.array([float(r[g]) for r in d]);assert np.isfinite(a).all();med=float(np.median(a))
 for r in d:r[g+'_group']='High' if float(r[g])>med else 'Low'
 thresholds.append(dict(gene=g,median_log2_FPKM_plus1=med,n_Low=sum(r[g+'_group']=='Low' for r in d),n_High=sum(r[g+'_group']=='High' for r in d),n_at_median=int(sum(a==med)),min=float(a.min()),max=float(a.max())))
for name,rows in [('patient_data_PRE73',d),('thresholds',thresholds)]:
 with (out/f'tables/{name}.csv').open('w',newline='') as h:
  wr=csv.DictWriter(h,fieldnames=list(rows[0]));wr.writeheader();wr.writerows(rows)
(out/'qa/input_manifest.json').write_text(json.dumps({str(p.relative_to(base)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [src,snapshot]},indent=2))
print(json.dumps(thresholds,indent=2))
