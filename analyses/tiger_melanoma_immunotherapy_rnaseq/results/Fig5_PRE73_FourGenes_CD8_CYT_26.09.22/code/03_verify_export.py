from pathlib import Path
import csv,json,math,hashlib,subprocess,shutil,ast
from collections import Counter
import numpy as np
from PIL import Image
from openpyxl import load_workbook

out=Path(__file__).resolve().parents[1];base=out.parent.parent
def read(p):
 with p.open() as h:return list(csv.DictReader(h))
checks=[]
def ck(name,ok):
 assert bool(ok),name
 checks.append(name)
def close(a,b):return math.isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-12)
# Reuse only the independently implemented numerical helper definitions, not the prior script's execution.
from numerical_helpers import rank,ibeta
def bh(ps):
 ps=np.array(ps);order=np.argsort(ps);q=np.empty(len(ps));q[order]=np.minimum(1,np.minimum.accumulate((ps[order]*len(ps)/np.arange(1,len(ps)+1))[::-1])[::-1]);return q
d=read(out/'tables/patient_data_PRE73.csv');cr=read(out/'tables/correlation_statistics.csv');gr=read(out/'tables/group_statistics.csv')
ar=lambda k:np.array([float(r[k]) for r in d])
ck('73 unique patients and samples',len(d)==len({r['sample_id'] for r in d})==len({r['patient_id'] for r in d})==73)
Z=np.column_stack([np.ones(73),rank(ar('age')),[r['sex']=='Male' for r in d],[r['therapy']=='combo' for r in d]])
ck('covariate matrix rank',np.linalg.matrix_rank(Z)==4)
res=read(out/'tables/rank_residuals.csv');independent=[]
for s in cr:
 m,v,model=s['metric'],s['endpoint'],s['model'];z=Z if model=='Adjusted' else np.ones((73,1))
 rx=rank(ar(m));ry=rank(ar(v));ex=rx-z@np.linalg.lstsq(z,rx,rcond=None)[0];ey=ry-z@np.linalg.lstsq(z,ry,rcond=None)[0]
 rho=float(np.corrcoef(ex,ey)[0,1]);df=73-z.shape[1]-1;t=rho*math.sqrt(df/(1-rho*rho));p=ibeta(df/2,.5,df/(df+t*t))
 key=m+' '+v+' '+model
 ck(key+' rho',close(rho,s['rho']));ck(key+' P',abs(math.log(p)-math.log(float(s['P'])))<1e-9);ck(key+' df',df==int(s['df']))
 rr={x['sample_id']:x for x in res if x['metric']==m and x['endpoint']==v and x['model']==model}
 for j,r in enumerate(d):ck(key+r['sample_id']+' residual',abs(ex[j]-float(rr[r['sample_id']]['x_residual']))<1e-10 and abs(ey[j]-float(rr[r['sample_id']]['y_residual']))<1e-10)
 independent.append(dict(kind='correlation',metric=m,endpoint=v,model=model,rho=rho,P=p))
for model in ['Adjusted','Unadjusted']:
 r=[s for s in cr if s['model']==model];ck(model+' BH8',np.allclose(bh([float(s['P']) for s in r]),[float(s['BH_q']) for s in r],atol=1e-12,rtol=1e-10))
for s in gr:
 m,v=s['metric'],s['endpoint'];lo=np.array([float(r[v]) for r in d if r[m+'_group']=='Low']);hi=np.array([float(r[v]) for r in d if r[m+'_group']=='High'])
 ck(m+v+' counts',len(lo)==int(s['n_Low']) and len(hi)==int(s['n_High']) and len(lo)+len(hi)==73)
 U=sum(float(h>l)+.5*float(h==l) for h in hi for l in lo);ct=Counter(np.r_[lo,hi]);N=73
 variance=len(lo)*len(hi)/12*(N+1-sum(t**3-t for t in ct.values())/(N*(N-1)))
 p=math.erfc(max(0,abs(U-len(lo)*len(hi)/2)-.5)/math.sqrt(variance*2))
 ck(m+v+' U',close(U,s['U_High']));ck(m+v+' P',close(p,s['P']));ck(m+v+' rank biserial',close(2*U/(len(lo)*len(hi))-1,s['rank_biserial']))
 for name,arr in [('Low',lo),('High',hi)]:
  for label,q in [('Q1',.25),('median',.5),('Q3',.75)]:ck(m+v+name+label,close(np.quantile(arr,q),s[name+'_'+label]))
 prism=read(out/'tables'/f'PRISM_{m}_{v}_groups.csv')
 for name,arr in [('Low',lo),('High',hi)]:ck(m+v+name+' Prism',np.allclose([float(r[name]) for r in prism if r[name]],arr,rtol=1e-12,atol=1e-12))
 pp=read(out/'qa'/f'{m}_{v}_scatter.csv');gg=read(out/'qa'/f'{m}_{v}_group.csv');xy=read(out/'tables'/f'PRISM_{m}_{v}_XY.csv')
 ck(m+v+' plot count',len(pp)==len(gg)==len(xy)==73)
 for i,r in enumerate(d):
  ck(m+v+r['sample_id']+' plot/Prism',pp[i]['sample_id']==gg[i]['sample_id']==xy[i]['sample_id']==r['sample_id'] and close(pp[i]['x'],r[m]) and close(pp[i]['y'],r[v]) and gg[i]['group']==r[m+'_group'] and close(gg[i]['value'],r[v]) and close(xy[i][m],r[m]) and close(xy[i][v],r[v]))
 independent.append(dict(kind='groups',metric=m,endpoint=v,U=U,P=p))
ck('group BH8',np.allclose(bh([float(s['P']) for s in gr]),[float(s['BH_q']) for s in gr],atol=1e-12,rtol=1e-10))
wb=load_workbook(out/'inputs/73Pre_snapshot.xlsx',data_only=True,read_only=True);sheet=wb['02_All_Patient_Data'];excel={row[1]:row for row in sheet.iter_rows(min_row=11,max_row=83,max_col=25,values_only=True)}
prior={r['sample_id']:r for r in read(base/'results/Fig5_PRE73_Balance_CD8_CYT_Correlations_26.09.22/tables/patient_data_PRE73.csv')}
metrics=['SARDH','PIPOX','GNMT','DMGDH']
for m in metrics:
 med=float(np.median(ar(m)))
 for r in d:ck(m+r['sample_id']+' median label',r[m+'_group']==('High' if float(r[m])>med else 'Low'))
for r in d:
 e=excel[r['sample_id']];p=prior[r['sample_id']]
 for k,i in [('SARDH',8),('PIPOX',9),('GNMT',10),('DMGDH',11)]:ck(k+r['sample_id']+' Excel',close(r[k],e[i]))
 for k in ['sample_id','patient_id','response','therapy','sex']:ck(k+r['sample_id']+' source',r[k]==p[k])
 for k in ['age','SARDH','PIPOX','GNMT','DMGDH','Balance','CD8_percent','CD8_fraction','CYT','GZMA_TPM','PRF1_TPM']:ck(k+r['sample_id']+' source',close(r[k],p[k]))
prior_heat=read(base/'results/Fig5_PRE73_Integrated_Immune_Heatmap_26.09.22/tables/partial_spearman_BH20.csv')
for row in cr:
 if row['model']!='Adjusted':continue
 previous=next(x for x in prior_heat if x['metric']==row['metric'] and x['endpoint']==row['endpoint'])
 ck(row['metric']+row['endpoint']+' original heatmap rho/P',close(row['rho'],previous['rho']) and close(row['P'],previous['P']))
for p,h in json.loads((out/'qa/input_manifest.json').read_text()).items():ck('unchanged source '+p,hashlib.sha256((base/p).read_bytes()).hexdigest()==h)
delivery=out.parent/'PPT';icc=Path('/System/Library/ColorSync/Profiles/sRGB Profile.icc').read_bytes();exports=read(out/'qa/figure_exports.csv')
for e in exports:
 src=out/'figures'/(e['stem']+'.png');im=Image.open(src).convert('RGB');im.save(src,dpi=(450,450),icc_profile=icc)
 dest=delivery/e['short'];shutil.copyfile(src,dest);subprocess.run(['xattr','-c',str(dest)],check=True);subprocess.run(['chflags','nohidden',str(dest)],check=True)
 saved=Image.open(dest);ck(e['short']+' RGB8',saved.mode=='RGB' and dest.read_bytes()[24]==8 and dest.read_bytes()[25]==2 and dest.read_bytes()[28]==0)
 ck(e['short']+' exact bytes',src.read_bytes()==dest.read_bytes());ck(e['short']+' short path',len(str(dest))<200)
 ck(e['short']+' no xattrs',not subprocess.check_output(['xattr',str(dest)],text=True).strip())
 # Contact thumbnails contain both endpoints for each metric without altering scientific files.
for kind in ['scatter','groups']:
 images=[]
 for m in metrics:
  im=Image.open(out/'figures'/f'{m}_{kind}_combined.png');im.thumbnail((1400,700));images.append(im)
 canvas=Image.new('RGB',(1400,sum(im.height for im in images)),'white');y=0
 for im in images:canvas.paste(im,(0,y));y+=im.height
 canvas.save(out/'qa'/f'{kind}_contact.png')
(out/'qa/verification.json').write_text(json.dumps(dict(checks_passed=len(checks),checks=checks,independent_results=independent,native_PPT_drag_tested=False),indent=2))
print(json.dumps(dict(checks_passed=len(checks),independent_results=independent),indent=2))
