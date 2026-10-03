from pathlib import Path
import csv,json,hashlib,math,subprocess,io,zipfile
import numpy as np
from PIL import Image,ImageCms
from pptx import Presentation
from pptx.util import Inches
out=Path(__file__).resolve().parents[1];base=out.parents[1];root=next(p for p in base.parents if (p/'PROJECT_HANDOFF.md').exists())
read=lambda p:list(csv.DictReader(p.open()))
d=read(out/'tables/patient_data_PRE73.csv');st=read(out/'tables/primary_logrank_results.csv')[0];cx=read(out/'tables/Cox_HR.csv')[0];km=read(out/'tables/KM_steps.csv');rt=read(out/'tables/risk_table.csv')
t=np.array([float(x['OS_days']) for x in d]);e=np.array([int(x['event']) for x in d]);x=np.array([int(x['group']=='High') for x in d]);bal=np.array([float(x['Balance']) for x in d]);checks=[]
def ck(name,ok):
 assert bool(ok),name
 checks.append({'check':name,'passed':True})
ck('73 unique and 29 events',len(d)==len({r['sample_id'] for r in d})==73 and sum(e)==29)
ck('median labels',np.array_equal(x,(bal>np.median(bal)).astype(int)))
times=sorted(set(t[e==1]));expected=0.;var=0.
for tt in times:
 risk=t>=tt;n=sum(risk);nlow=sum(risk&(x==0));deaths=sum((t==tt)&(e==1));expected+=deaths*nlow/n
 if n>1:var+=deaths*(nlow/n)*(1-nlow/n)*(n-deaths)/(n-1)
chi=(sum(e[x==0])-expected)**2/var;p=math.erfc(math.sqrt(chi/2))
ck('logrank_expected',abs(expected-float(st['expected_Low']))<1e-12);ck('logrank_statistic',abs(chi-float(st['chi_square']))<1e-12);ck('logrank_P',abs(p-float(st['P_logrank_nominal']))<1e-12)
for group,flag in [('Low',0),('High',1)]:
 ix=x==flag;s=1.;calc={0:(1.,sum(ix),0,0)}
 for tt in sorted(set(t[ix])):
  n=sum(ix&(t>=tt));ev=sum(ix&(t==tt)&(e==1));cen=sum(ix&(t==tt)&(e==0));s*=1-ev/n;calc[tt]=(s,n,ev,cen)
 for row in [r for r in km if r['group']==group]:
  val=calc[float(row['time'])];ck(group+'_KM_'+row['time'],abs(val[0]-float(row['survival']))<1e-12 and val[1]==int(row['n_risk']) and val[2]==int(row['n_event']) and val[3]==int(row['n_censor']))
 for row in [r for r in rt if r['group']==group]:ck(group+'_risk_'+row['time'],sum(ix&(t>=float(row['time'])))==int(row['n_risk']))
 gs=next(r for r in read(out/'tables/group_summary.csv') if r['group']==group);below=[tt for tt,val in calc.items() if val[0]<=.5];median=min(below) if below else None
 ck(group+'_median',gs['median_OS_days']=='NA' if median is None else float(gs['median_OS_days'])==median)
def efron(beta):
 ll=0.;u=0.;info=0.;w=np.exp(beta*x)
 for tt in times:
  risk=t>=tt;die=(t==tt)&(e==1);nd=int(sum(die));S0=sum(w[risk]);S1=sum((w*x)[risk]);S2=sum((w*x*x)[risk]);E0=sum(w[die]);E1=sum((w*x)[die]);E2=sum((w*x*x)[die]);ll+=beta*sum(x[die]);u+=sum(x[die])
  for j in range(nd):
   frac=j/nd;den=S0-frac*E0;a=S1-frac*E1;b=S2-frac*E2;ll-=math.log(den);u-=a/den;info+=b/den-(a/den)**2
 return ll,u,info
beta=0.
for _ in range(100):
 ll,u,info=efron(beta);step=u/info;beta+=step
 if abs(step)<1e-13:break
ll,u,info=efron(beta);se=1/math.sqrt(info);hr=math.exp(beta);ci=[math.exp(beta-1.959963984540054*se),math.exp(beta+1.959963984540054*se)]
for name,val in [('HR',hr),('coef',beta),('se',se),('CI_low',ci[0]),('CI_high',ci[1]),('Wald_P',math.erfc(abs(beta/se)/math.sqrt(2)))]:ck('Cox_'+name,abs(float(cx[name])-val)<1e-9)
ck('plotted_KM_equals_source',read(out/'qa/plotted_KM.csv')==km)
ck('plotted_risk_equals_source',read(out/'qa/plotted_risk.csv')==rt)
pr=read(out/'tables/PRISM_OS_Balance.csv');actual=sorted((float(r['Time_days']),'Low' if r['Balance_Low']!='' else 'High',int(r['Balance_Low'] if r['Balance_Low']!='' else r['Balance_High'])) for r in pr);want=sorted((float(r['OS_days']),r['group'],int(r['event'])) for r in d);ck('Prism_all73',actual==want)
for name,m in json.loads((out/'qa/input_manifest.json').read_text()).items():ck(name+'_source_unchanged',hashlib.sha256((root/m['relative_path']).read_bytes()).hexdigest()==m['sha256'])
png=out/'figures/Balance_HighLow_OS.png';im=Image.open(png).convert('RGB');pixels=np.asarray(im).copy();icc=Path('/System/Library/ColorSync/Profiles/sRGB Profile.icc').read_bytes();im.save(png,dpi=(450,450),icc_profile=icc)
transport=out/'figures/PPT_insert/B.png';transport.write_bytes(png.read_bytes());delivery=base/'results/PPT/BOS.png';delivery.parent.mkdir(exist_ok=True);delivery.write_bytes(png.read_bytes())
for target in [transport,delivery]:
 subprocess.run(['xattr','-c',str(target)],check=True);subprocess.run(['chflags','nohidden',str(target)],check=True);si=Image.open(target)
 ck(target.name+'_identical_pixels',np.array_equal(pixels,np.asarray(si)));ck(target.name+'_RGB8sRGB',si.mode=='RGB' and si.info['icc_profile']==icc and target.read_bytes()[24]==8 and target.read_bytes()[25]==2 and target.read_bytes()[28]==0);ck(target.name+'_xattr_clear',not subprocess.check_output(['xattr',str(target)],text=True).strip())
prs=Presentation();prs.slide_width=Inches(7);prs.slide_height=Inches(7.3);slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(delivery),0,0,width=prs.slide_width,height=prs.slide_height);scratch=out/'qa/scratch.pptx';prs.save(scratch)
with zipfile.ZipFile(scratch) as z:ck('scratch_integrity',z.testzip() is None)
reopen=Presentation(scratch);pic=reopen.slides[0].shapes[0];ck('scratch_embedded_bytes',pic.image.blob==delivery.read_bytes());Image.open(io.BytesIO(pic.image.blob)).resize((1150,round(1150*7.3/7)),Image.Resampling.LANCZOS).save(out/'qa/scratch_render.png')
print('Independent checks:',len(checks),'PASS; P=',p,'HR=',hr)
(out/'qa/verification.json').write_text(json.dumps({'checks':checks,'independent_logrank_P':p,'independent_HR':hr,'native_PowerPoint_drag_verified':False,'short_delivery':str(delivery),'short_delivery_path_chars':len(str(delivery))},indent=2))
