from pathlib import Path
import csv,json,hashlib,math
import numpy as np
from openpyxl import load_workbook
out=Path(__file__).resolve().parents[1];base=out.parents[1];root=next(p for p in base.parents if (p/'PROJECT_HANDOFF.md').exists())
sources={'excel':base/'73Pre.xlsx','original73':base/'results/tables/analysis_data_PRE.csv','previous_groups':base/'results/Fig5_PRE73_Delta_HighLow_OS_26.09.20/tables/patient_classifications_PRE73.csv','previous_logrank':base/'results/Fig5_PRE73_Delta_HighLow_OS_26.09.20/tables/primary_logrank_results.csv','previous_cox':base/'results/Fig5_PRE73_Delta_HighLow_OS_26.09.20/tables/Cox_HR.csv','balance_correlation':base/'results/Fig5_PRE73_Balance_CD8_CYT_Correlations_26.09.22/tables/patient_data_PRE73.csv'}
manifest={k:dict(relative_path=str(p.relative_to(root)),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for k,p in sources.items()}
read=lambda p:list(csv.DictReader(p.open()))
original=read(sources['original73']);previous=read(sources['previous_groups']);balance=read(sources['balance_correlation'])
orig={r['sample_id']:r for r in original};prev={r['sample_id']:r for r in previous};bc={r['sample_id']:r for r in balance}
s=load_workbook(sources['excel'],data_only=True).active;rows=list(s.iter_rows(min_row=11,max_row=83,values_only=True))
assert len(orig)==len(prev)==len(bc)==len(rows)==73 and set(orig)==set(prev)==set(bc)=={r[1] for r in rows}
assert len({r[0] for r in rows})==73
cut=float(np.median([r[21] for r in rows]));data=[]
for r in rows:
 sid=r[1];a=orig[sid];v=prev[sid];c=bc[sid];g='High' if r[21]>cut else 'Low'
 assert r[0]==a['patient_name']==v['patient_name']==c['patient_id'] and a['timepoint']==v['timepoint']=='PRE'
 assert r[2]==a['response_group']==v['response_group']==c['response'] and r[3]==a['therapy_short']==v['therapy_short']==c['therapy']
 assert r[6]==float(a['os_days'])==float(v['os_days']) and r[7]==int(a['event'])==int(v['event'])==int(a['vital_status']=='Dead')
 assert r[21]==float(c['Balance']) and abs(r[21]-float(v['Delta'])*math.sqrt(73/72))<1e-12 and g==v['delta_group']
 assert r[6]>0 and r[7] in [0,1]
 data.append(dict(sample_id=sid,patient_id=r[0],Balance=r[21],group=g,OS_days=r[6],event=r[7],response=r[2],therapy=r[3],age=r[4],sex=r[5]))
assert sum(r['event'] for r in data)==29 and sum(r['group']=='High' for r in data)==36
with (out/'tables/patient_data_PRE73.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=data[0]);w.writeheader();w.writerows(data)
(out/'qa/input_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False))
for k,p in sources.items():assert hashlib.sha256(p.read_bytes()).hexdigest()==manifest[k]['sha256']
print('73 patients matched; all Balance groups and OS match original and priorDelta. Cutoff:',cut)
