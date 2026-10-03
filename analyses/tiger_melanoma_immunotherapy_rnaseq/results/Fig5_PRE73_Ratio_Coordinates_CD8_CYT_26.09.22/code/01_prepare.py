from pathlib import Path
import csv,json,hashlib,math
import numpy as np
from openpyxl import load_workbook

out=Path(__file__).resolve().parents[1];base=out.parent.parent
for sub in ['tables','figures','qa']:(out/sub).mkdir(exist_ok=True)
def read(p):
 with p.open() as h:return list(csv.DictReader(h))
sources=[out/'inputs/73Pre_snapshot.xlsx',base/'results/tables/analysis_data_PRE.csv',base/'results/Fig5_PRE73_Balance_CD8_CYT_Correlations_26.09.22/tables/patient_data_PRE73.csv']
immune=read(sources[2]);old={r['sample_id']:r for r in read(sources[1])}
wb=load_workbook(sources[0],read_only=True,data_only=True);wf=load_workbook(sources[0],read_only=True,data_only=False)
sh=wb['02_All_Patient_Data'];sf=wf['02_All_Patient_Data']
headers=[sh.cell(10,c).value for c in range(1,26)]
excel=[dict(zip(headers,row)) for row in sh.iter_rows(min_row=11,max_row=83,max_col=25,values_only=True)]
assert len(excel)==len({r['sample_id'] for r in excel})==len({r['patient_id'] for r in excel})==73
ex={r['sample_id']:r for r in excel};assert set(ex)==set(old)=={r['sample_id'] for r in immune}
genes=['SARDH','PIPOX','GNMT','DMGDH'];G=np.array([[float(r[g]) for g in genes] for r in excel]);Z=(G-G.mean(0))/G.std(0,ddof=0)
checks=[]
for i,e in enumerate(excel):
 assert all(math.isclose(e['z'+g],Z[i,j],rel_tol=1e-11,abs_tol=1e-12) for j,g in enumerate(genes))
 assert e['Pro_sum']>0
 calculated={'Deg_Pro_ratio':(e['SARDH']+e['PIPOX'])/(e['GNMT']+e['DMGDH']),
             'Deg_coord':(Z[i,0]+Z[i,1])/math.sqrt(2),'Pro_coord':(Z[i,2]+Z[i,3])/math.sqrt(2)}
 for k,v in calculated.items():assert math.isclose(e[k],v,rel_tol=1e-10,abs_tol=1e-12)
 for new,prev in [('Deg_coord','Degradation_score'),('Pro_coord','Production_score')]:
  assert math.isclose(e[new],float(old[e['sample_id']][prev])*math.sqrt(2*73/72),rel_tol=1e-10,abs_tol=1e-12)
 checks.append(e['sample_id'])
rows=[]
for r in immune:
 e=ex[r['sample_id']]
 for k in ['patient_id','response','therapy','sex']:assert r[k]==e[k]
 assert float(r['age'])==e['age']
 for k in genes+['Balance']:assert math.isclose(float(r[k]),e[k],rel_tol=1e-11,abs_tol=1e-12)
 n=dict(r)
 for k in ['Deg_sum','Pro_sum','Deg_Pro_ratio','Deg_coord','Pro_coord']:n[k]=e[k]
 for k in ['Degradation_score','Production_score']:n[k]=old[r['sample_id']][k]
 assert old[r['sample_id']]['timepoint']=='PRE'
 rows.append(n)
metrics=['Deg_Pro_ratio','Deg_coord','Pro_coord'];thresholds=[]
for m in metrics:
 vals=np.array([r[m] for r in rows]);assert np.isfinite(vals).all()
 med=float(np.median(vals))
 for r in rows:r[m+'_group']='High' if r[m]>med else 'Low'
 thresholds.append(dict(metric=m,median=med,n_Low=sum(r[m+'_group']=='Low' for r in rows),n_High=sum(r[m+'_group']=='High' for r in rows),min=float(vals.min()),max=float(vals.max())))
 if m!='Deg_Pro_ratio':
  prev='Degradation_score' if m=='Deg_coord' else 'Production_score';pmed=np.median([float(r[prev]) for r in rows])
  assert all(r[m+'_group']==('High' if float(r[prev])>pmed else 'Low') for r in rows)
for name,data in [('patient_data_PRE73',rows),('thresholds',thresholds)]:
 with (out/f'tables/{name}.csv').open('w',newline='') as h:
  w=csv.DictWriter(h,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
(out/'qa/input_manifest.json').write_text(json.dumps({str(p.relative_to(base)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},indent=2))
(out/'qa/excel_formulas.json').write_text(json.dumps({sf.cell(r,c).coordinate:sf.cell(r,c).value for r in range(11,84) for c in [19,20,21]},indent=2))
(out/'qa/preparation.json').write_text(json.dumps(dict(patient_count=73,validated_ids=checks,coordinate_scale_to_old=math.sqrt(2*73/72),thresholds=thresholds,min_ratio_denominator=min(r['Pro_sum'] for r in rows)),indent=2))
print(json.dumps(thresholds,indent=2))
