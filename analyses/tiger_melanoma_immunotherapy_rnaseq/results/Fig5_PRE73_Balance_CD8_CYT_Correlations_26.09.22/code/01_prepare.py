from pathlib import Path
import csv,json,hashlib,math
import numpy as np
from openpyxl import load_workbook
out=Path(__file__).resolve().parents[1];base=out.parents[1]
root=next(p for p in base.parents if (p/'PROJECT_HANDOFF.md').exists())
heat=next(root.glob('*/CRC&NSCLC-TCGA/Fig6_TCGA_validation/results/Fig5_Production_Degradation_CD8_CYT_Correlations_26.09.18/tables/patient_source_data.csv'))
sources={'excel':base/'73Pre.xlsx','immune':base/'results/Fig5_PRE73_CD8_CYT_OS_26.09.21/tables/patient_data_PRE73.csv','delta':base/'results/Fig5_PRE73_DegMinusProd_Response_26.09.20/tables/patient_data_PRE73.csv','current_heatmap':heat}
hashes={k:{'relative_path':str(v.relative_to(root)),'sha256':hashlib.sha256(v.read_bytes()).hexdigest()} for k,v in sources.items()}
checks=[]
def ck(name,ok):
 assert ok,name
 checks.append({'check':name,'passed':True})
def csvread(p):return list(csv.DictReader(p.open()))
def write(name,rows):
 with (out/'tables'/name).open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
s=load_workbook(sources['excel'],data_only=True).active
f=load_workbook(sources['excel'],data_only=False).active
rows=list(s.iter_rows(min_row=11,max_row=83,values_only=True))
ck('exact73unique',len(rows)==73 and len({r[0] for r in rows})==73 and len({r[1] for r in rows})==73)
ck('noadditionalpatientrows',all(s.cell(i,1).value is None for i in range(84,s.max_row+1)))
immune=csvread(sources['immune']);delta=csvread(sources['delta']);h=[r for r in csvread(heat) if r['cohort']=='Melanoma']
for name,dd in [('immune',immune),('delta',delta),('heatmap',h)]:
 ck(name+'_unique_roster',len(dd)==73 and len({x['sample_id'] for x in dd})==73 and {x['sample_id'] for x in dd}=={r[1] for r in rows})
immune={r['sample_id']:r for r in immune};delta={r['sample_id']:r for r in delta};h={r['sample_id']:r for r in h}
a=np.array([r[8:12] for r in rows],float);zs=(a-a.mean(0))/a.std(0,ddof=0)
ck('populationSD_zscores',np.max(abs(zs-np.array([r[12:16] for r in rows],float)))<1e-12)
joined=[]
for j,r in enumerate(rows):
 sid=r[1];u=immune[sid];d=delta[sid];v=h[sid]
 ck(sid+'_metadata',r[0]==u['patient_id']==d['patient_id']==v['patient_id'] and r[2]==u['response']==d['response_group'] and r[3]==u['therapy']==d['therapy']==v['context'] and r[4]==float(u['age'])==float(d['age'])==float(v['age']) and r[5].upper()==u['sex'].upper()==d['sex'].upper()==v['sex'].upper() and u['timepoint']==d['timepoint']=='PRE')
 ck(sid+'_genes',all(abs(r[c]-float(d[g]))<1e-12 and abs(r[c]-float(v[g]))<1e-12 for c,g in zip(range(8,12),['SARDH','PIPOX','GNMT','DMGDH'])))
 bal=(zs[j,0]+zs[j,1]-zs[j,2]-zs[j,3])/2
 ck(sid+'_Balance',abs(r[21]-bal)<1e-12 and abs(r[21]-float(d['Deg_minus_Prod'])*math.sqrt(73/72))<1e-12 and f.cell(j+11,22).value==f'=(M{j+11}+N{j+11}-O{j+11}-P{j+11})/2')
 cd8=float(u['CD8_fraction']);cyt=float(u['CYT'])
 ck(sid+'_immune',abs(cd8-float(v['CD8']))<1e-12 and abs(cyt-float(v['CYT']))<1e-12 and 0<=cd8<=1 and abs(cyt-math.log2(math.sqrt(float(u['GZMA_TPM'])*float(u['PRF1_TPM']))))<1e-12)
 row=dict(sample_id=sid,patient_id=r[0],response=r[2],therapy=r[3],age=r[4],sex=r[5],Balance=r[21],Delta_sampleSD=float(d['Deg_minus_Prod']),CD8_fraction=cd8,CD8_percent=100*cd8,CYT=cyt,GZMA_TPM=float(u['GZMA_TPM']),PRF1_TPM=float(u['PRF1_TPM']))
 row.update(dict(zip(['SARDH','PIPOX','GNMT','DMGDH'],r[8:12])));joined.append(row)
ck('complete_numeric',all(np.isfinite([x[k] for k in ['Balance','CD8_percent','CYT','age']]).all() for x in joined))
write('patient_data_PRE73.csv',joined)
for v,stem in [('CD8_percent','CD8'),('CYT','CYT')]:write('PRISM_Balance_'+stem+'.csv',[dict(sample_id=r['sample_id'],patient_id=r['patient_id'],Balance=r['Balance'],**{v:r[v]}) for r in joined])
for k,p in sources.items():ck(k+'_source_unchanged',hashlib.sha256(p.read_bytes()).hexdigest()==hashes[k]['sha256'])
(out/'qa/input_manifest.json').write_text(json.dumps(hashes,indent=2,ensure_ascii=False))
(out/'qa/preparation_checks.json').write_text(json.dumps(checks,indent=2))
print('Preparation:',len(checks),'checks PASS;73 complete patients')
