from pathlib import Path
import csv,json,math,hashlib,struct,zlib,zipfile,os,subprocess
import numpy as np
from PIL import Image,ImageCms
from pptx import Presentation
from pptx.util import Inches
out=Path(__file__).resolve().parents[1];root=next(p for p in out.parents if (p/'PROJECT_HANDOFF.md').exists())
checks=[]
def ck(name,ok):
 assert bool(ok),name
 checks.append({'check':name,'passed':True})
def read(p):return list(csv.DictReader(p.open()))
def rank(x):
 a=np.asarray(x);order=np.argsort(a,kind='stable');r=np.empty(len(a),float);i=0
 while i<len(a):
  j=i+1
  while j<len(a) and a[order[j]]==a[order[i]]:j+=1
  r[order[i:j]]=(i+1+j)/2;i=j
 return r
# Continued fraction for the regularized incomplete beta; independent of R's pt.
def betacf(a,b,x):
 qab=a+b;qap=a+1;qam=a-1;c=1.;d=1.-qab*x/qap
 if abs(d)<1e-300:d=1e-300
 d=1/d;h=d
 for m in range(1,401):
  m2=2*m;aa=m*(b-m)*x/((qam+m2)*(a+m2));d=1+aa*d
  if abs(d)<1e-300:d=1e-300
  c=1+aa/c
  if abs(c)<1e-300:c=1e-300
  d=1/d;h*=d*c;aa=-(a+m)*(qab+m)*x/((a+m2)*(qap+m2));d=1+aa*d
  if abs(d)<1e-300:d=1e-300
  c=1+aa/c
  if abs(c)<1e-300:c=1e-300
  d=1/d;delta=d*c;h*=delta
  if abs(delta-1)<3e-14:return h
 raise AssertionError('beta failed convergence')
def ibeta(a,b,x):
 if x==0:return 0.
 if x==1:return 1.
 bt=math.exp(math.lgamma(a+b)-math.lgamma(a)-math.lgamma(b)+a*math.log(x)+b*math.log1p(-x))
 return bt*betacf(a,b,x)/a if x<(a+1)/(a+b+2) else 1-bt*betacf(b,a,1-x)/b
D=read(out/'tables/patient_data_PRE73.csv');R=read(out/'tables/correlation_statistics.csv')
a=lambda k:np.array([float(x[k]) for x in D])
Z=np.column_stack([np.ones(73),rank(a('age')),[x['sex']=='Male' for x in D],[x['therapy']=='combo' for x in D]])
ck('matrix_rank4',np.linalg.matrix_rank(Z)==4)
res=read(out/'tables/rank_residuals.csv');ind=[]
for row in R:
 v=row['endpoint'];adj=row['model']=='Adjusted';z=Z if adj else np.ones((73,1));rx=rank(a('Balance'));ry=rank(a(v))
 ex=rx-z@np.linalg.lstsq(z,rx,rcond=None)[0];ey=ry-z@np.linalg.lstsq(z,ry,rcond=None)[0];rho=float(np.corrcoef(ex,ey)[0,1]);df=73-z.shape[1]-1;t=rho*math.sqrt(df/(1-rho*rho));p=ibeta(df/2,.5,df/(df+t*t))
 stem=v+'_'+row['model'];ck(stem+'_rho',abs(rho-float(row['rho']))<1e-12);ck(stem+'_P',abs(math.log(p)-math.log(float(row['P'])))<1e-10);ck(stem+'_df',df==int(row['df']))
 records={x['sample_id']:x for x in res if x['endpoint']==v and x['model']==row['model']}
 for j,x in enumerate(D):ck(stem+'_'+x['sample_id']+'_residuals',abs(ex[j]-float(records[x['sample_id']]['Balance_residual']))<1e-10 and abs(ey[j]-float(records[x['sample_id']]['outcome_residual']))<1e-10)
 ind.append(dict(endpoint=v,model=row['model'],rho=rho,P=p))
for m in ['Adjusted','Unadjusted']:
 r=[x for x in R if x['model']==m];ps=np.array([float(x['P']) for x in r]);order=np.argsort(ps);q=np.empty(2);q[order]=np.minimum(1,np.minimum.accumulate((ps[order]*2/np.arange(1,3))[::-1])[::-1])
 ck(m+'_BH2',max(abs(q-np.array([float(x['BH_q']) for x in r])))<1e-12)
for v,stem in [('CD8_percent','Balance_CD8'),('CYT','Balance_CYT')]:
 pp=read(out/'qa'/f'{stem}_plotted.csv');ck(stem+'_all73',len(pp)==73)
 ck(stem+'_coordinates',all(p['sample_id']==d['sample_id'] and abs(float(p['x'])-float(d['Balance']))<1e-12 and abs(float(p['y'])-float(d[v]))<1e-12 for p,d in zip(pp,D)))
 for x in read(out/'tables'/f'PRISM_{stem}.csv'):ck(stem+'_prism_'+x['sample_id'],any(d['sample_id']==x['sample_id'] and float(d['Balance'])==float(x['Balance']) and float(d[v])==float(x[v]) for d in D))
# Standalone scientific PNGs and pixel-identical transport copies.
profile=Path('/System/Library/ColorSync/Profiles/sRGB Profile.icc');icc=profile.read_bytes() if profile.exists() else ImageCms.ImageCmsProfile(ImageCms.createProfile('sRGB')).tobytes()
render_dir=out/'qa';image_checks=[]
prs=Presentation();prs.slide_width=Inches(13);prs.slide_height=Inches(5.8)
for stem in ['Balance_CD8','Balance_CYT','Balance_Combined']:
 src=out/'figures'/f'{stem}.png';im=Image.open(src).convert('RGB');pixels=np.asarray(im).copy();im.save(src,dpi=(450,450),icc_profile=icc)
 transport=out/'figures/PPT_insert'/f'PPT_INSERT_{stem}.png';transport.write_bytes(src.read_bytes())
 subprocess.run(['xattr','-c',str(transport)],check=True);subprocess.run(['chflags','nohidden',str(transport)],check=True)
 saved=Image.open(transport);ck(stem+'_pixels',np.array_equal(pixels,np.asarray(saved)));ck(stem+'_RGB8',saved.mode=='RGB' and transport.read_bytes()[24]==8 and transport.read_bytes()[25]==2 and transport.read_bytes()[28]==0);ck(stem+'_icc',bool(saved.info.get('icc_profile')));ck(stem+'_dpi',max(abs(x-450) for x in saved.info['dpi'])<.1);ck(stem+'_xattrs',not subprocess.check_output(['xattr',str(transport)],text=True).strip())
 # Three scratch slides: original plot aspect preserved at full height, centered.
 slide=prs.slides.add_slide(prs.slide_layouts[6]);ratio=im.width/im.height;h=prs.slide_height;w=int(h*ratio);left=int((prs.slide_width-w)/2)
 slide.shapes.add_picture(str(transport),left,0,width=w,height=h)
 image_checks.append({'stem':stem,'dimensions':im.size,'dpi':saved.info['dpi'],'profile':ImageCms.getProfileDescription(ImageCms.ImageCmsProfile(str(profile))) if profile.exists() else 'sRGB','sha256':hashlib.sha256(transport.read_bytes()).hexdigest()})
scratch=render_dir/'scratch_embedding.pptx';prs.save(scratch)
with zipfile.ZipFile(scratch) as z:ck('scratch_zip',z.testzip() is None)
reopened=Presentation(scratch)
import io
for j,slide in enumerate(reopened.slides):
 pic=slide.shapes[0];stem=image_checks[j]['stem'];original=(out/'figures/PPT_insert'/f'PPT_INSERT_{stem}.png').read_bytes();ck(stem+'_scratch_exact_embedded_bytes',pic.image.blob==original)
 canvas=Image.new('RGB',(1800,round(1800*reopened.slide_height/reopened.slide_width)),'white')
 px=Image.open(io.BytesIO(pic.image.blob)).convert('RGB');scale=1800/reopened.slide_width;box=(round(pic.left*scale),round(pic.top*scale));size=(round(pic.width*scale),round(pic.height*scale));canvas.paste(px.resize(size,Image.Resampling.LANCZOS),box);canvas.save(render_dir/f'{stem}_scratch_render.png')
# Independent macOS decoder checks output dimensions.
for row in image_checks:
 p=out/'figures/PPT_insert'/f"PPT_INSERT_{row['stem']}.png"
 sp=subprocess.check_output(['sips','-g','pixelWidth','-g','pixelHeight',str(p)],text=True)
 ck(row['stem']+'_sips_dimensions',f"pixelWidth: {row['dimensions'][0]}" in sp and f"pixelHeight: {row['dimensions'][1]}" in sp)
for k,r in json.loads((out/'qa/input_manifest.json').read_text()).items():ck(k+'_hash_unchanged',hashlib.sha256((root/r['relative_path']).read_bytes()).hexdigest()==r['sha256'])
# Keep the author-confirmed short-path delivery convention for subsequent reruns.
shortdir=out.parents[1]/'results/PPT';shortdir.mkdir(exist_ok=True)
for stem,shortname in [('Balance_CD8','BC.png'),('Balance_CYT','BY.png'),('Balance_Combined','BCY.png')]:
 target=shortdir/shortname;target.write_bytes((out/'figures'/f'{stem}.png').read_bytes())
 subprocess.run(['xattr','-c',str(target)],check=True);subprocess.run(['chflags','nohidden',str(target)],check=True)
 assert target.read_bytes()==(out/'figures'/f'{stem}.png').read_bytes()
(out/'qa/independent_verification.json').write_text(json.dumps({'checks':checks,'independent_statistics':ind,'image_checks':image_checks,'scratch_render_method':'Picture geometry and embedded PNG read from saved/reopened PPTX, composed with Pillow. Native PowerPoint rendering/drag behavior not certified.'},indent=2))
print('Independent verification:',len(checks),'checks PASS');print(ind)
