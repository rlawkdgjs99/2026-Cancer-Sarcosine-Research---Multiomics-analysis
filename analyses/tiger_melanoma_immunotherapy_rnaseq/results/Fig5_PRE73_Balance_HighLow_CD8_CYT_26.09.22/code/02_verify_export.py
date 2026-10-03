from pathlib import Path
import csv, json, hashlib, math, subprocess, shutil
from collections import Counter
import numpy as np
from openpyxl import load_workbook
from PIL import Image

out=Path(__file__).resolve().parent.parent
base=out.parent.parent
checks=[]
def check(name,condition):
    if not condition: raise AssertionError(name)
    checks.append(name)
def read(path):
    with path.open() as h:return list(csv.DictReader(h))
def close(a,b):return math.isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-12)
d=read(out/'tables/patient_data_PRE73.csv')
prior=read(out.parent/'Fig5_PRE73_Balance_CD8_CYT_Correlations_26.09.22/tables/patient_data_PRE73.csv')
os={r['sample_id']:r for r in read(out.parent/'Fig5_PRE73_Balance_HighLow_OS_26.09.22/tables/patient_data_PRE73.csv')}
check('73 unique patients and samples',len(d)==len({r['sample_id'] for r in d})==len({r['patient_id'] for r in d})==73)
wb=load_workbook(base/'73Pre.xlsx',read_only=True,data_only=True)
sheet=wb['02_All_Patient_Data']
excel={sheet.cell(r,2).value:sheet.cell(r,22).value for r in range(11,84)}
check('Excel ID set',set(excel)=={r['sample_id'] for r in d})
cut=float(np.median([float(r['Balance']) for r in d]))
for r,p in zip(d,prior):
    for k in p:
        check('source '+r['sample_id']+' '+k,r[k]==p[k] if k in ['sample_id','patient_id','response','therapy','sex'] else close(r[k],p[k]))
    check('Excel Balance '+r['sample_id'],close(r['Balance'],excel[r['sample_id']]))
    check('OS labels '+r['sample_id'],r['Balance_group']==os[r['sample_id']]['group'])
    check('median label '+r['sample_id'],r['Balance_group']==('High' if float(r['Balance'])>cut else 'Low'))
stats=read(out/'tables/comparison_statistics.csv'); pvalues=[];calculated=[]
for st in stats:
    v=st['outcome'];lo=np.array([float(r[v]) for r in d if r['Balance_group']=='Low']);hi=np.array([float(r[v]) for r in d if r['Balance_group']=='High'])
    check(v+' n',len(lo)==37 and len(hi)==36)
    U=sum(float(h>l)+.5*float(h==l) for h in hi for l in lo)
    counts=Counter(np.r_[lo,hi]);N=len(lo)+len(hi)
    variance=len(lo)*len(hi)/12*(N+1-sum(t**3-t for t in counts.values())/(N*(N-1)))
    z=(abs(U-len(lo)*len(hi)/2)-.5)/math.sqrt(variance)
    p=math.erfc(z/math.sqrt(2));pvalues.append(p)
    check(v+' independent U',close(U,st['U_High']))
    check(v+' independent P',close(p,st['p']))
    check(v+' rank biserial',close(2*U/(len(lo)*len(hi))-1,st['rank_biserial']))
    for group,arr in [('Low',lo),('High',hi)]:
        for label,q in [('Q1',.25),('median',.5),('Q3',.75)]:check(v+group+label,close(np.quantile(arr,q),st[group+'_'+label]))
    plot=read(out/('qa/plotted_'+v+'.csv'))
    check(v+' plotted points',len(plot)==73)
    for r,s in zip(d,plot):check(v+' plot '+r['sample_id'],r['sample_id']==s['sample_id'] and r['Balance_group']==s['group'] and close(r[v],s['value']))
    prism=read(out/('tables/PRISM_'+v+'.csv'))
    for group,arr in [('Low',lo),('High',hi)]:check(v+group+' Prism',np.allclose([float(r[group]) for r in prism if r[group]],arr,rtol=1e-12,atol=1e-12))
    calculated.append(dict(outcome=v,U=U,tie_variance=variance,z=z,p=p))
order=np.argsort(pvalues);q=np.zeros(2);running=1.
for i in range(1,-1,-1):
    j=order[i];running=min(running,pvalues[j]*2/(i+1));q[j]=running
for i,st in enumerate(stats):check(st['outcome']+' BH',close(q[i],st['BH_q']))
for path,sha in json.loads((out/'qa/input_manifest.json').read_text()).items():check('unchanged '+path,hashlib.sha256((base/path).read_bytes()).hexdigest()==sha)
delivery=out.parent/'PPT';delivery.mkdir(exist_ok=True)
for stem,short in [('CD8_percent','BGC.png'),('CYT','BGY.png'),('Combined','BGCY.png')]:
    src=out/('figures/Balance_HighLow_'+stem+'.png')
    im=Image.open(src);im.load()
    if im.mode!='RGB':im=im.convert('RGB')
    icc=Path('/System/Library/ColorSync/Profiles/sRGB Profile.icc').read_bytes()
    im.save(src,dpi=(450,450),icc_profile=icc)
    dest=delivery/short;shutil.copyfile(src,dest)
    subprocess.run(['xattr','-c',str(dest)],check=True);subprocess.run(['chflags','nohidden',str(dest)],check=True)
    check(short+' bytes',src.read_bytes()==dest.read_bytes())
    check(short+' RGB',Image.open(dest).mode=='RGB')
    check(short+' short path',len(str(dest))<200)
im=Image.open(delivery/'BGCY.png');im.thumbnail((1500,900));im.save(out/'qa/combined_preview.png')
(out/'qa/verification.json').write_text(json.dumps(dict(checks_passed=len(checks),checks=checks,independent_results=calculated,native_PowerPoint_drag_tested=False),indent=2))
print(json.dumps(dict(checks_passed=len(checks),independent_results=calculated),indent=2))
