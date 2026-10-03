from pathlib import Path
import csv,hashlib,shutil,json
from openpyxl import load_workbook
o=Path(__file__).resolve().parent.parent
b=o.parent.parent
for d in ['code','inputs','tables','figures','qa']: (o/d).mkdir(parents=True,exist_ok=True)
if not (o/'inputs/16Matched_snapshot.xlsx').exists():
 shutil.copyfile(b/'16Matched.xlsx',o/'inputs/16Matched_snapshot.xlsx')
w=load_workbook(o/'inputs/16Matched_snapshot.xlsx',read_only=True,data_only=True).worksheets[0]
headers=[c.value for c in w[3]]
rows=[dict(zip(headers,r)) for r in w.iter_rows(min_row=4,values_only=True) if r[0] is not None]
assert len(rows)==16
fields=[h for h in headers if h is not None]
with (o/'inputs/matched16_source.csv').open('w',newline='') as f:
 wr=csv.DictWriter(f,fields,extrasaction='ignore');wr.writeheader();wr.writerows(rows)
for src,dest in [
 (b/'results/manuscript_figures/Fig2d_CYT_candidate_26.08.23/SourceData_CYT_all91_sensitivity.csv','CYT_all91.csv'),
 (b/'results/Fig5_PRE73_FourGenes_CD8_CYT_26.09.22/tables/patient_data_PRE73.csv','prior_PRE73.csv'),
 (b/'results/TIGER_PRE73_Degradation_HighLow_Deconvolution_7methods_26.09.02/raw_deconvolution/deconvolution_quanTIseq.csv','prior_quanTIseq_PRE73.csv')]:shutil.copyfile(src,o/'inputs'/dest)
(o/'ANALYSIS_PLAN.md').write_text('''# Matched16 PRE–EDT plan (fixed before outcome testing)
Same 16 patients in author 16Matched.xlsx, R9/NR7, each PRE+EDT; all therapies, no cutpoints/exclusions. Six endpoints: CD8_percent, CYT, SARDH, PIPOX, GNMT, DMGDH.
Use original gene log2(FPKM+1); CD8% from existing IOBR2.2.3 quanTIseq settings (linear FPKM, tumor TRUE, arrays FALSE, scale_mrna TRUE; same local reference). Recompute both matched timepoints, check PRE against existing73 result. CYT=log2(sqrt(GZMA_TPM*PRF1_TPM)), using full transcriptome FPKM-to-TPM per sample, verify stored all91 CYT. Verify original clinical PRE/EDT IDs and all workbook gene values against expression matrix.
Primary question: delta=EDT−PRE differs R vs NR? Two-sided unpaired Mann–Whitney on16patient deltas, normal approximation/tie+continuity correction; BH6 acrosssixendpoints. Direct rank/distribution comparison; not causal therapy effect or covariate-adjusted interaction.
Additional: overall paired PRE–EDT two-sided Wilcoxon signed-rank, normal approximation, continuity correction, digits.rank12, BH6; R- and NR-stratified paired tests BH12 across6×2. Distribution of paired differences assumed symmetric for signed-rank location interpretation. Small n means limited power. Zeros removed only from signed-rank ranking as method requires, retained in patient tables/plots.
Secondary cross-sectional R vs NR at PRE and at EDT: Mann–Whitney same settings, BH12 across6×2, reported separately. Do not conclude between-group differences from one within-group significant result and other nonsignificant. No pooled32sample independent comparison. No changes to prior73figures/statistics.
Plots: paired PRE–EDT lines overall and byR/NR, deltaRvsNR dots/medianIQR. CD8 changes are percentage points; genes change log2(FPKM+1), CYT change log2geometricmeanTPM. Show BHq and preserve rawP, individualIDs, alloutcomes, fullstats/PrismCSV. Supplementary exact sign-flip signed-rank and exact rank-label permutation checks may be used as small-sample sensitivity, not to replace primary method by significance.
''')
paths=[b/'16Matched.xlsx',b/'Melanoma-PRJEB23709_ClinicalData.tsv',b.parent/'사용데이터_모음/Source_Input/TIGER_PRJEB23709_expression_FPKM_gene_by_sample.csv',b/'reference_data/IOBR_v2.2.3_data-v1.0/quantiseq_data.rda']
# Unicode path resolution for sibling directory.
import unicodedata
used=next(p for p in b.parent.iterdir() if unicodedata.normalize('NFC',p.name)=='사용데이터_모음')
paths[2]=used/'Source_Input/TIGER_PRJEB23709_expression_FPKM_gene_by_sample.csv'
for p in paths:assert p.exists(),p
(o/'inputs/source_manifest.json').write_text(json.dumps({str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},ensure_ascii=False,indent=2))
(o/'inputs/expression_path.txt').write_text(str(paths[2]))

print(o)
print((b/'Melanoma-PRJEB23709_ClinicalData.tsv').read_text()[:1300])
