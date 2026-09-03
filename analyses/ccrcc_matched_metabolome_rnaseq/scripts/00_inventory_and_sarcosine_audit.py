#!/usr/bin/env python3
"""Read-only inventory of the supplied ccRCC multi-omics files.

The script never modifies source files.  It records workbook structure, selected
cell previews, metabolite-name hits, and the expression-matrix schema so that
the downstream sample-matching decision is auditable.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

import openpyxl


PROJECT = Path(__file__).resolve().parents[1]
SOURCE_DATA = PROJECT / "논문_SourceData"
SUP_DATA = PROJECT / "논문_SupData"
RNA_FILE = PROJECT / "RNAseq_Data" / "bulkRNA_matrix_TPM.csv"
OUT_DIR = PROJECT / "results" / "00_input_audit_26.09.02"

SEARCH_TERMS = {
    "sarcosine": re.compile(r"\bsarcosine\b", re.I),
    "n_methylglycine": re.compile(r"n[\s_-]*methylglycine|methylglycine", re.I),
    "dimethylglycine": re.compile(r"dimethylglycine", re.I),
    "glycine": re.compile(r"\bglycine\b", re.I),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def display_value(value: object, limit: int = 160) -> object:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value).replace("\r", " ").replace("\n", " ")
    return text if len(text) <= limit else text[: limit - 3] + "..."


def audit_workbook(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    record = {
        "path": str(path.relative_to(PROJECT)),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "sheets": [],
        "term_hits": [],
    }
    for sheet in workbook.worksheets:
        preview = []
        nonempty = 0
        hits = []
        for row_index, row in enumerate(sheet.iter_rows(), start=1):
            preview_row = []
            for column_index, cell in enumerate(row, start=1):
                value = cell.value
                if value is not None:
                    nonempty += 1
                    if isinstance(value, str):
                        for term, pattern in SEARCH_TERMS.items():
                            if pattern.search(value):
                                hits.append(
                                    {
                                        "term": term,
                                        "sheet": sheet.title,
                                        "cell": cell.coordinate,
                                        "value": display_value(value, 500),
                                    }
                                )
                if row_index <= 8 and column_index <= 14:
                    preview_row.append(display_value(value))
            if preview_row:
                preview.append(preview_row)
        record["sheets"].append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_cells": nonempty,
                "preview_A1_N8": preview,
                "term_hit_count": len(hits),
            }
        )
        record["term_hits"].extend(hits)
    workbook.close()
    return record


def audit_rna_csv(path: Path) -> dict:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        first_rows = []
        row_count = 0
        widths = {len(header)}
        gene_names = []
        numeric_min = None
        numeric_max = None
        zero_count = 0
        numeric_count = 0
        for row in reader:
            row_count += 1
            widths.add(len(row))
            if row:
                gene_names.append(row[0])
            if row_count <= 5:
                first_rows.append(row[:8])
            for raw in row[1:]:
                if raw == "":
                    continue
                value = float(raw)
                numeric_count += 1
                zero_count += int(value == 0)
                numeric_min = value if numeric_min is None else min(numeric_min, value)
                numeric_max = value if numeric_max is None else max(numeric_max, value)
    sample_ids = header[1:]
    suffix_counts = {
        "tumor_T": sum(sample.endswith("_T") for sample in sample_ids),
        "normal_N": sum(sample.endswith("_N") for sample in sample_ids),
        "other": sum(not (sample.endswith("_T") or sample.endswith("_N")) for sample in sample_ids),
    }
    patient_ids = [re.sub(r"_[TN]$", "", sample) for sample in sample_ids]
    return {
        "path": str(path.relative_to(PROJECT)),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "rows_features": row_count,
        "columns_including_feature_id": len(header),
        "row_widths_observed": sorted(widths),
        "sample_count": len(sample_ids),
        "unique_sample_count": len(set(sample_ids)),
        "unique_patient_count": len(set(patient_ids)),
        "suffix_counts": suffix_counts,
        "sample_ids": sample_ids,
        "duplicate_sample_ids": sorted({x for x in sample_ids if sample_ids.count(x) > 1}),
        "duplicate_gene_names_count": len(gene_names) - len(set(gene_names)),
        "first_rows_first_8_columns": first_rows,
        "numeric_min": numeric_min,
        "numeric_max": numeric_max,
        "numeric_zero_fraction": zero_count / numeric_count if numeric_count else None,
    }


def main() -> None:
    xlsx_files = sorted(SOURCE_DATA.glob("*.xlsx")) + sorted(SUP_DATA.glob("*.xlsx"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    audit = {
        "project": str(PROJECT),
        "workbooks": [audit_workbook(path) for path in xlsx_files],
        "rna_csv": audit_rna_csv(RNA_FILE),
    }
    output = OUT_DIR / "input_inventory_and_sarcosine_hits.json"
    output.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")

    with (OUT_DIR / "sarcosine_term_hits.tsv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["workbook", "term", "sheet", "cell", "value"],
            delimiter="\t",
        )
        writer.writeheader()
        for workbook in audit["workbooks"]:
            for hit in workbook["term_hits"]:
                writer.writerow({"workbook": workbook["path"], **hit})

    print(json.dumps({
        "workbooks": len(audit["workbooks"]),
        "sheets": sum(len(w["sheets"]) for w in audit["workbooks"]),
        "term_hits": sum(len(w["term_hits"]) for w in audit["workbooks"]),
        "rna": {
            key: audit["rna_csv"][key]
            for key in [
                "rows_features", "sample_count", "unique_patient_count",
                "suffix_counts", "row_widths_observed", "numeric_min",
                "numeric_max", "numeric_zero_fraction",
            ]
        },
        "output": str(output),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
